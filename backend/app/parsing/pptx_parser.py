"""PPTX parser."""

from __future__ import annotations

from pathlib import Path

from pptx import Presentation

from backend.app.schemas.document import DocumentStructure, Section


def extract_pptx(path: Path | str) -> DocumentStructure:
    prs = Presentation(str(path))
    sections: list[Section] = []
    full: list[str] = []

    for idx, slide in enumerate(prs.slides, start=1):
        texts: list[str] = []
        title = f"Slide {idx}"
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue
            t = (shape.text or "").strip()
            if not t:
                continue
            if shape == slide.shapes.title and t:
                title = t
            texts.append(t)
            full.append(t)
        sections.append(Section(heading=title, level=1, text="\n".join(texts), page_start=idx))

    return DocumentStructure(
        title=sections[0].heading if sections else None,
        page_count=len(prs.slides),
        sections=sections,
        full_text="\n\n".join(full),
        parser_route="python_pptx",
    )
